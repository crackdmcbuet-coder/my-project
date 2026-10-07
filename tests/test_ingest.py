import pytest
from docx import Document

from ingest.extract import NoTextLayerError, Unit, chunk_units, extract_units
from ingest.structure import entry_id, process_book, verify

from .conftest import FakeClient

PAGE = "সাপ দেখা\nএটি একটি নমুনা ব্যাখ্যা।"


def chunk_of(*texts):
    return chunk_units([Unit(f"পৃষ্ঠা {i}", t) for i, t in enumerate(texts, 1)], max_chars=10_000)[0]


def test_verify_accepts_exact_copy_even_across_whitespace_changes():
    chunk = chunk_of("সাপ  দেখা\nএটি একটি\nনমুনা ব্যাখ্যা।")
    assert verify("সাপ দেখা", "এটি একটি নমুনা ব্যাখ্যা।", chunk)


def test_verify_rejects_reworded_or_invented_text():
    chunk = chunk_of(PAGE)
    assert not verify("সাপ দেখা", "এটি একটি ভিন্ন ব্যাখ্যা।", chunk)
    assert not verify("পানি দেখা", "এটি একটি নমুনা ব্যাখ্যা।", chunk)


def test_verify_works_across_page_markers():
    chunk = chunk_of("সাপ দেখা এটি একটি", "নমুনা ব্যাখ্যা।")
    assert "[পৃষ্ঠা 2]" in chunk.marked and "[পৃষ্ঠা" not in chunk.plain
    assert verify("সাপ দেখা", "এটি একটি নমুনা ব্যাখ্যা।", chunk)


def test_chunking_overlaps_by_one_unit_and_covers_everything():
    units = [Unit(f"পৃষ্ঠা {i}", "ক" * 40) for i in range(1, 8)]
    chunks = chunk_units(units, max_chars=100, overlap_units=1)
    assert len(chunks) > 1
    assert chunks[0].labels[-1] == chunks[1].labels[0]
    assert {l for c in chunks for l in c.labels} == {u.label for u in units}


def test_process_book_splits_verified_from_needs_review_and_dedupes_overlap():
    chunk = chunk_of(PAGE)
    payload = {"entries": [
        {"dream": "সাপ দেখা", "interpretation": "এটি একটি নমুনা ব্যাখ্যা।", "label": "[পৃষ্ঠা 1]"},
        {"dream": "সাপ দেখা", "interpretation": "এটি একটি নমুনা ব্যাখ্যা।", "label": "পৃষ্ঠা 1"},  # dup
        {"dream": "আগুন দেখা", "interpretation": "বানানো ব্যাখ্যা", "label": "পৃষ্ঠা 1"},
    ]}
    entries, review, t_in, t_out = process_book(FakeClient(payload), "m", "বই", [chunk], lambda *_: None)
    assert [e["source"] for e in entries] == ["বই, পৃষ্ঠা 1"]
    assert entries[0]["id"] == entry_id("সাপ দেখা", "এটি একটি নমুনা ব্যাখ্যা।")
    assert [r["dream"] for r in review] == ["আগুন দেখা"] and review[0]["verified"] is False
    assert (t_in, t_out) == (100, 20)


def test_extract_docx_includes_tables_in_order(tmp_path):
    doc = Document()
    doc.add_paragraph("ভূমিকা")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "সাপ দেখা", "ব্যাখ্যা"
    doc.add_paragraph("শেষ")
    path = tmp_path / "b.docx"
    doc.save(path)
    assert [u.text for u in extract_units(path)] == ["ভূমিকা", "সাপ দেখা | ব্যাখ্যা", "শেষ"]


def test_extract_html_skips_scripts(tmp_path):
    path = tmp_path / "b.html"
    path.write_text("<html><script>x=1</script><body><p>সাপ দেখা</p><p>ব্যাখ্যা</p></body></html>", encoding="utf-8")
    assert [u.text for u in extract_units(path)] == ["সাপ দেখা", "ব্যাখ্যা"]


def test_scanned_pdf_gets_a_clear_error(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(200, 200)
    path = tmp_path / "scan.pdf"
    with open(path, "wb") as f:
        writer.write(f)
    with pytest.raises(NoTextLayerError):
        extract_units(path)


def test_unsupported_type(tmp_path):
    path = tmp_path / "b.txt"
    path.write_text("x")
    with pytest.raises(ValueError):
        extract_units(path)
